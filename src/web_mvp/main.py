"""Interactive FastHTML citizen surface for JanaVani.

This application is an independent web interface. Drafting is delegated to the
existing authenticated AI adapter; generated review documents are templates
and are never represented as official submissions.
"""
from __future__ import annotations

from html import escape
from io import BytesIO
import os
import re
from typing import Tuple
from xml.sax.saxutils import escape as xml_escape

from docx import Document
from fasthtml.common import (
    A,
    Br,
    Button,
    Container,
    Div,
    Form,
    H1,
    H2,
    H3,
    H4,
    Hr,
    Hidden,
    Label,
    Link,
    P,
    Radio,
    Span,
    Style,
    Textarea,
    Titled,
    fast_app,
    rt,
)
from fastapi.responses import Response
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

from src.adapters.web_client import JanavaniAIWebClient


app, rt = fast_app(
    hdrs=(
        Link(rel="preconnect", href="https://cdn.jsdelivr.net"),
        Style("""
            body { padding: 2rem 0; background-color: #f8f9fa; }
            .container { max-width: 900px; }
            .header-banner { text-align: center; margin-bottom: 3rem; }
            .card { background: white; padding: 2rem; border-radius: 8px;
                    box-shadow: 0 4px 6px rgba(0,0,0,0.05); margin-bottom: 2rem; }
            .triangle-badge { display: inline-block; padding: 0.25rem 0.75rem;
                    border-radius: 50px; font-size: 0.85rem; font-weight: bold;
                    background-color: #fff1d6; color: #7a4b00; }
        """),
    )
)


def build_objection_document(
    bill_code: str,
    comments: str,
    format_choice: str,
) -> Tuple[bytes, str, str]:
    """Build a citizen review-request template, not a legal opinion."""
    safe_code = bill_code.strip()
    if not re.fullmatch(r"[A-Za-z0-9_-]{1,80}", safe_code):
        raise ValueError("Invalid bill or record reference.")
    selected_format = format_choice.strip().upper()
    if selected_format not in {"PDF", "DOCX"}:
        raise ValueError("Select PDF or DOCX.")

    content_lines = [
        "Citizen Request for Review — Draft",
        f"Reference: {safe_code}",
        "",
        "To the appropriate public authority,",
        "",
        "I request review of the public record identified above and ask that the "
        "authority provide the applicable text, decision-making basis, and a "
        "method for submitting comments through the proper official channel.",
        "",
        "Citizen observations (provided by the user):",
        comments.strip() or "No additional observations supplied.",
        "",
        "This is a citizen-prepared template. It is not legal advice, proof that "
        "the referenced record is authentic, or evidence of submission to any authority.",
    ]

    if selected_format == "DOCX":
        document = Document()
        document.add_heading(content_lines[0], 0)
        for line in content_lines[1:]:
            document.add_paragraph(line)
        output = BytesIO()
        document.save(output)
        return (
            output.getvalue(),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            f"citizen_review_draft_{safe_code}.docx",
        )

    output = BytesIO()
    pdf = SimpleDocTemplate(output)
    style = getSampleStyleSheet()["BodyText"]
    story = []
    for line in content_lines:
        story.append(Paragraph(xml_escape(line) if line else "&#160;", style))
        story.append(Spacer(1, 8))
    pdf.build(story)
    return output.getvalue(), "application/pdf", f"citizen_review_draft_{safe_code}.pdf"


@rt("/")
def home_page():
    """Render the citizen action hub and draft form."""
    return Titled(
        "🇮🇳 Janavani — Citizen Action Hub",
        Container(
            Div(
                H1("JANAVANI"),
                P(
                    "Turn civic concerns into structured drafts through "
                    "privacy-conscious workflows."
                ),
                style="text-align: center; margin-bottom: 3rem;",
            ),
            Div(
                H3("📝 Prepare a civic grievance draft"),
                P(
                    "Describe the issue in your preferred language. The text is "
                    "sent to the configured Janavani drafting service; avoid "
                    "including unnecessary personal or sensitive information."
                ),
                Form(action="/submit-issue", method="post")(
                    Textarea(
                        name="citizen_input",
                        placeholder=(
                            "Describe the civic issue, where it occurred, and "
                            "what action you are requesting."
                        ),
                        rows=5,
                        required=True,
                    ),
                    Button("Prepare Draft", type="submit", cls="button-primary"),
                ),
                cls="card",
            ),
            Div(
                H3("⚖️ Legislative and public-record review"),
                P(
                    "Use the review-request template to organize observations "
                    "about a public record. Verify the source text and legal "
                    "position against authoritative material before relying on it."
                ),
                Div(
                    H4("Illustrative record: BILL-2026-KL-04"),
                    Span("Illustrative — not independently verified", cls="triangle-badge"),
                    P(
                        "This example is a workflow demonstration, not a finding "
                        "that a real bill exists or violates constitutional rights."
                    ),
                    A(
                        "Prepare a review-request draft",
                        href="/bill-review/BILL-2026-KL-04",
                        cls="button secondary",
                    ),
                    cls="card",
                ),
                cls="card",
            ),
        ),
    )


@rt("/submit-issue")
def submit_issue(citizen_input: str):
    """Request a draft from the configured AI service and show honest status."""
    client = JanavaniAIWebClient()
    result = client.request_document_draft(citizen_input)
    if not isinstance(result, dict):
        return Titled(
            "Draft service unavailable",
            Container(
                Div(
                    H2("Draft could not be prepared"),
                    P(
                        "The configured drafting service did not return a usable "
                        "response. Your text has not been submitted to a government authority."
                    ),
                    A("Return to JanaVani", href="/"),
                    cls="card",
                )
            ),
        )

    draft_text = result.get("draft")
    model_result = result.get("result")
    if isinstance(model_result, dict):
        choices = model_result.get("choices", [])
        if choices and isinstance(choices[0], dict):
            message = choices[0].get("message", {})
            if isinstance(message, dict) and isinstance(message.get("content"), str):
                draft_text = message["content"]

    is_degraded = result.get("status") != "available"
    if not isinstance(draft_text, str) or not draft_text.strip():
        draft_text = result.get("message", "No readable draft was returned.")

    response_items = [
        H2("Civic document draft"),
        P(
            "DRAFT ONLY — this page does not register, file, or submit a complaint."
        ),
    ]
    if is_degraded:
        response_items.append(
            P(
                "The AI drafting service was unavailable or degraded. The displayed "
                "text may be the original input; it is not a verified legal analysis."
            )
        )
    response_items.extend(
        [
            Div(draft_text, style="white-space: pre-wrap; overflow-wrap: anywhere;"),
            Br(),
            A("Return to JanaVani", href="/"),
        ]
    )
    return Titled(
        "Civic draft — review before use",
        Container(Div(*response_items, cls="card")),
    )


@rt("/bill-review/{bill_code}")
def review_bill(bill_code: str):
    """Render a user-observation form for an illustrative public record."""
    return Titled(
        "Public-record review request",
        Container(
            Div(
                H2(f"Review request: {bill_code}"),
                P(
                    "This template does not independently verify the record or "
                    "make a legal finding. Check the official text and source."
                ),
                Hr(),
                Form(action="/dispatch-objection", method="post")(
                    Hidden(name="bill_code", value=bill_code),
                    H4("Your observations (optional)"),
                    Textarea(
                        name="comments",
                        placeholder=(
                            "Describe the source, passage, concern, and requested "
                            "clarification. Avoid unnecessary sensitive information."
                        ),
                        rows=5,
                    ),
                    H4("Document format"),
                    Label(Radio(name="format_choice", value="PDF", checked=True), " PDF"),
                    Label(Radio(name="format_choice", value="DOCX"), " Editable DOCX"),
                    Br(),
                    Button("Generate review-request draft", type="submit"),
                ),
                A("Cancel and return", href="/"),
                cls="card",
            )
        ),
    )


@rt("/dispatch-objection")
def dispatch_objection(
    bill_code: str,
    comments: str = "",
    format_choice: str = "PDF",
):
    """Generate a local review-request template; never submit it externally."""
    try:
        file_bytes, media_type, filename = build_objection_document(
            bill_code, comments, format_choice
        )
    except ValueError as error:
        return Titled(
            "Invalid review request",
            Container(
                Div(
                    P(str(error)),
                    A("Return to the form", href=f"/bill-review/{escape(bill_code)}"),
                    cls="card",
                )
            ),
        )

    return Response(
        content=file_bytes,
        media_type=media_type,
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


if __name__ == "__main__":
    import uvicorn

    uvicorn.run(app, host="0.0.0.0", port=int(os.getenv("PORT", "8080")))
