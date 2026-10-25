"""In-memory rendering of anonymous civic drafts; no blob store or filesystem persistence."""
from __future__ import annotations

from io import BytesIO
from xml.sax.saxutils import escape as xml_escape

from src.capabilities.anonymous_civic_drafting import AnonymousCivicDraft


def render_anonymous_draft(draft: AnonymousCivicDraft, document_format: str) -> tuple[bytes, str, str]:
    """Render a transient draft to PDF or DOCX bytes and return (bytes, media type, extension)."""
    kind = document_format.strip().casefold()
    if kind == "docx":
        from docx import Document

        buffer = BytesIO()
        document = Document()
        document.add_heading(draft.title, level=1)
        document.add_heading(draft.subject, level=2)
        for paragraph in draft.body.split("\n\n"):
            document.add_paragraph(paragraph)
        document.add_paragraph(
            "DRAFT ONLY — review the facts and applicable process before sending. "
            "Janavani has not submitted this document."
        )
        document.save(buffer)
        return (
            buffer.getvalue(),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "docx",
        )
    if kind == "pdf":
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.lib.units import mm
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

        buffer = BytesIO()
        document = SimpleDocTemplate(
            buffer, pagesize=A4, rightMargin=20 * mm, leftMargin=20 * mm,
            topMargin=18 * mm, bottomMargin=18 * mm,
            title=draft.title, author="Janavani — citizen-prepared draft",
        )
        styles = getSampleStyleSheet()
        story = [Paragraph(xml_escape(draft.title), styles["Title"]), Spacer(1, 8)]
        story.append(Paragraph(xml_escape(draft.subject), styles["Heading2"]))
        story.append(Spacer(1, 8))
        for paragraph in draft.body.split("\n\n"):
            story.append(Paragraph(xml_escape(paragraph).replace("\n", "<br/>"), styles["BodyText"]))
            story.append(Spacer(1, 7))
        story.append(Spacer(1, 8))
        story.append(Paragraph(
            "<b>DRAFT ONLY.</b> Review facts and applicable process before sending. "
            "Janavani has not submitted this document.",
            styles["Italic"],
        ))
        document.build(story)
        return buffer.getvalue(), "application/pdf", "pdf"
    raise ValueError("Choose PDF or DOCX")
